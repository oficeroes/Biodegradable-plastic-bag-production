from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).parent


def set_font(style, size, bold=False):
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Microsoft YaHei")
    fonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def new_doc(title, lead):
    d = Document()
    section = d.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = d.styles
    set_font(styles["Normal"], 10.2)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Normal"].paragraph_format.line_spacing = 1.2
    set_font(styles["Title"], 17, True)
    styles["Title"].paragraph_format.space_after = Pt(9)
    title_ppr = styles["Title"].element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    set_font(styles["Heading 1"], 12.3, True)
    styles["Heading 1"].paragraph_format.space_before = Pt(11)
    styles["Heading 1"].paragraph_format.space_after = Pt(5)
    set_font(styles["Heading 2"], 10.6, True)
    styles["Heading 2"].paragraph_format.space_before = Pt(8)
    styles["Heading 2"].paragraph_format.space_after = Pt(3)
    p = d.add_paragraph(style="Title")
    p.add_run(title)
    d.add_paragraph(lead)
    return d


def para(d, text, bold_prefix=None):
    p = d.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def bullet(d, text):
    p = d.add_paragraph(style="Normal")
    p.style = d.styles["Normal"]
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.add_run("• ")
    p.add_run(text)
    return p


def heading(d, text, level=1):
    d.add_paragraph(text, style=f"Heading {level}")


def table(d, headers, rows, widths=None):
    t = d.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for idx, h in enumerate(headers):
        cell = t.rows[0].cells[idx]
        cell.text = h
        if widths:
            cell.width = Inches(widths[idx])
        set_cell_shading(cell, "E9ECEF")
        for run in cell.paragraphs[0].runs:
            run.bold = True
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row in rows:
        cells = t.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
            if widths:
                cells[idx].width = Inches(widths[idx])
    for row in t.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margin(cell, 80, 80, 100, 100)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                for r in p.runs:
                    r.font.name = "Microsoft YaHei"
                    r.font.size = Pt(8.8)
                    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            set_cell_border(cell)
    d.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def set_cell_border(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "D9D9D9")
        borders.append(e)


def set_cell_margin(cell, top, bottom, left, right):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for side, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        tag = OxmlElement(f"w:{side}")
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")
        margins.append(tag)


def small(d, text):
    p = d.add_paragraph(text)
    for r in p.runs:
        r.font.size = Pt(8.8)
    return p


def lab_proposal():
    d = new_doc(
        "小龙虾壳提取壳聚糖及成膜实验申请书",
        "申请事项：在学校实验室、教师监督和既有化学品管理制度下，以 10.00 g 干小龙虾壳开展一批提取验证，并用其中 1.00 g 壳聚糖制作一张非食品接触薄膜。申请范围不包括家中提取、餐具使用或购物袋承重宣传。",
    )
    table(d, ["申请信息", "填写处"], [
        ("申请人 / 班级", "____________________________"),
        ("指导教师 / 实验室", "____________________________"),
        ("拟定日期 / 场地", "____________________________"),
    ], [1.55, 5.45])

    heading(d, "一 研究依据与申请边界")
    para(d, "本方案以 Zullo 等 2026 年对 Faxonius limosus 小龙虾壳的研究为主要工艺依据。其干壳到壳聚糖收率为 15 ± 1%（3 个独立重复）；10.00 g 干壳仅可预计约 1.5 ± 0.1 g 壳聚糖，这不是本实验保证产量。论文未证明袋体承重、湿态耐用或食品接触安全。我国常见克氏原螯虾与论文物种不同，必须记录原料来源与批次。")
    para(d, "申请采用论文的固液比、温度和时间，另把 NaOH 的百分比定义写成项目拟定条件：去蛋白液为 3.5% w/v（每 100 mL 最终溶液含 3.50 g NaOH），脱乙酰液为 50% w/w（溶液总质量的一半为 NaOH）。原文只写 3.5% 和 50%，没有声明 w/v 或 w/w；这两项定义及其与原文的可比性须由指导教师核定。")

    heading(d, "二 10.00 g 干壳批次投料表")
    table(d, ["材料 / 试剂", "精确投料或计算规则", "来源与说明"], [
        ("小龙虾干壳粉", "10.00 g", "清洗、干燥后称量；湿壳所需重量未知。"),
        ("预配 1.000 M HCl", "150.0 mL；含 HCl 0.150 mol = 5.47 g 纯 HCl", "论文 1 g 干壳 : 15 mL 酸液。勿按估计体积自行稀释浓盐酸。"),
        ("预配 3.5% w/v NaOH", "39.0 × m₁ mL；m₁ = 去矿后干固体实称克数", "含 NaOH 1.365 × m₁ g；纯 NaOH 克重以教师确认的 w/v 定义为前提。"),
        ("预配 50% w/w NaOH", "10.0 × m₂ mL；m₂ = 去蛋白后干几丁质实称克数", "先量取再称实际溶液质量 M；其中纯 NaOH = 0.500 × M g。原文浓度基准未声明。"),
        ("去离子水", "洗涤量按终点决定，逐次记录；先备 ≥2 L", "2 L 只是备料量，不是固定消耗或排放量。"),
        ("成膜材料", "本批若得到 ≥1.00 g 壳聚糖，可另用 1.00 g 做薄膜", "成膜用量、酸液及甘油见第五节；不足 1.00 g 不补称不明材料。"),
    ], [1.55, 2.75, 2.7])
    para(d, "计算例：若实测 m₁ = 5.00 g，去蛋白液需 195.0 mL，按 3.5% w/v 含 NaOH 6.825 g；若 m₂ = 2.50 g，脱乙酰液需 25.0 mL。示例质量不代表本批预期收率。50% 溶液的准确克重需由本批实际称量得到，不能用未经测量的密度推断。")

    heading(d, "三 场地设备与人员要求")
    bullet(d, "场地：可用通风柜、洗眼器和安全冲淋；实验室负责人确认耐浓热碱的容器、搅拌器、水浴、温度监测、过滤方式及化学废液桶。不得使用铝器皿，酸反应容器不得密闭。")
    bullet(d, "人员：浓酸碱的取用、50% NaOH 溶液准备、加热和转移只由具备培训资格的人员按本校 SOP 执行或现场监督。学生先阅读各试剂 SDS，佩戴护目镜、实验服、适用手套；热碱操作另用面屏和防溅防护。")
    bullet(d, "器材：0.001 g 分析天平、量筒、pH 计或经确认的试纸、温度计、45°C 与 80°C 干燥设备、PTFE 包覆搅拌子、耐化学腐蚀的过滤/洗涤器具、标签及独立酸碱废液容器。")

    heading(d, "四 操作步骤与停止点")
    table(d, ["步骤", "操作与记录", "停止或放行条件"], [
        ("1 原料", "将壳与肌肉组织分离，流水刷洗；70°C 水中加热 15–20 min，反复洗净；80°C 干燥 1 h，冷却后研磨并称 10.00 g。记录物种、来源及壳批次。", "有甲壳类过敏或哮喘的参与者不进入壳粉操作。粉碎扬尘时停止并改进控制。"),
        ("2 去矿", "在通风柜中将 10.00 g 干壳粉与 150.0 mL 1.000 M HCl 接触，室温搅拌；CO₂ 气泡停止且液相仍呈酸性时结束。过滤，去离子水洗至末次洗液近中性；45°C 干燥 18 h，冷却称 m₁。", "容器全程敞口且有足够余量。异常喷溅、过热、pH 不再呈酸性时停止，由教师决定处理。"),
        ("3 去蛋白", "教师核定 3.5% w/v NaOH 后，取 39.0 × m₁ mL 溶液；42 ± 1°C 搅拌 3.5 h。冷却、过滤、洗至末次洗液近中性。为精确计算下一步投料，项目增设 45°C 干至恒重并称 m₂。", "若试剂标签浓度不是 3.5% w/v，需重新计算并获批。中间体干燥为本项目补充步骤，记录是否影响质量。"),
        ("4 脱乙酰", "教师确认 50% w/w NaOH 储液和适配器具后，取 10.0 × m₂ mL，称储液实重 M；约 100°C 沸水浴、搅拌 6 h。先在安全位置冷却，再过滤和洗至近中性；室温干燥至质量稳定，称成品 m₃。", "这是最高风险步骤。通风、冲淋、监护或耐腐蚀器具任一缺失时不开展；不得在家复制。"),
        ("5 质量核对", "记录 m₁、m₂、m₃ 和每次液体实耗量；计算干壳收率 100 × m₃ / 10.00%。测成品溶解性、灰分/残留蛋白、脱乙酰度和黏度的能力由教师决定。", "仅以‘能成膜’不能证明成品纯度、食品安全或降解性能。"),
    ], [1.0, 4.15, 1.85])
    para(d, "本文的‘末次洗液近中性’拟作为项目内部判据：连续两次末次洗液 pH 约 6–8，并记录原始值；这不是论文给出的完整纯度标准。若洗液仍强酸或强碱，继续按实验室 SOP 洗涤，不能仅凭干燥结束处理。")

    heading(d, "五 成膜验证与记录")
    para(d, "仅在本批得到至少 1.00 g 可溶壳聚糖后进行。项目标准化小试：壳聚糖 1.00 g、由实验室预配的 1.00% w/w 醋酸水溶液 100.00 g、甘油 0.30 g；室温搅拌 24 h、静置脱泡 2 h，全部倒入 15 × 15 cm 水平模具，65°C 初干 5 h，必要时继续干至质量稳定；再在记录温湿度的环境中放置 48 h。此配方借鉴商品壳聚糖复合膜研究并改为质量计量，不是 Zullo 原文的 1 g 壳聚糖 + 100 mL 水 + 2 mL 醋酸无甘油配方。")
    para(d, "每张膜至少测五点厚度，记录外观、能否完整脱模、干态折叠及接触水后的形态。只做材料验证，不宣称为可承重购物袋。若本批壳聚糖不足 1.00 g，则只记录提取和材料表征结果。")

    heading(d, "六 危害控制与废弃物")
    bullet(d, "热浓碱主要造成深部皮肤和眼睛灼伤；酸会喷溅并与壳中碳酸盐放出 CO₂。不得将酸碱在皮肤上互相中和，接触时立即持续用流动清水冲洗并按实验室应急程序就医。")
    bullet(d, "酸液、碱液、热碱洗液和含蛋白清洗液分容器收集，标注组分、浓度、日期和责任人；由学校按照现行制度处置，不直接倒入下水道。")
    bullet(d, "不进行丙酮或 H₂O₂/HCl 脱色，不做食品接触或人体使用。任何浓度基准、批量、加热设备或步骤改变都需教师重新批准。")

    heading(d, "七 教师审批")
    para(d, "请指导教师在开始实验前确认：① 3.5% NaOH 与 50% NaOH 的质量基准；② 加热浓碱的器皿、通风与个人防护；③ 50% 碱液的领取/配制责任；④ 洗涤终点和废液去向；⑤ 学生可独立操作的步骤。")
    table(d, ["审批结论", "填写处"], [
        ("意见", "□同意按方案实施  □修改后再审  □暂不同意"),
        ("需修改事项", "________________________________________________________"),
        ("指导教师签名 / 日期", "________________________________________________________"),
    ], [1.65, 5.35])
    heading(d, "主要依据")
    small(d, "Zullo R 等. New Active Biopolymers and Chitosan-Based Films from Non-Native Crayfish Shell. Molecules, 2026. DOI: 10.3390/molecules31162819，方法 3.3–3.4。原文存于项目 papers/01_Zullo_2026_Crayfish_Films.pdf。")
    small(d, "Hromiš N 等. Influence of Chitosan Extraction Process from Invasive Crayfish Shells. Gels, 2026. DOI: 10.3390/gels12080664。用于不同提取路线的风险与收率比较，不作为本申请的主操作 SOP。")
    small(d, "Jiménez-Regalado 等. Preparation and Physicochemical Properties of Modified Corn Starch-Chitosan Biodegradable Films. Polymers, 2021. DOI: 10.3390/polym13244431。用于甘油比例和纯壳聚糖膜的可行性参考。实验前须另核对学校 SDS 和 SOP。")
    d.save(ROOT / "学校实验室_虾壳提取壳聚糖_申请方案.docx")


def bought_chitosan_protocol():
    d = new_doc(
        "商品壳聚糖制作小袋样品操作流程",
        "用途：直接购买有规格证明的壳聚糖，制成约 11 × 13.5 cm 的小型袋体原型，验证薄膜、接缝和提手是否可承受实际提举。本流程不包含虾壳提取，所得袋体在性能与安全测试前不用于装食品，也不称为合格购物袋。",
    )
    heading(d, "一 配方来源和项目设定")
    para(d, "2021 年改性玉米淀粉/壳聚糖研究的纯壳聚糖组使用 1% w/v 商品壳聚糖、1% v/v 醋酸溶液，以及相当于聚合物质量 30% 的甘油；论文向 15 × 15 cm 模具浇铸约 100 mL 成膜液，65°C 干燥 5 h，随后约 25°C、54% RH 调湿 48 h。该组实测膜厚 49.07 ± 2.59 μm（n = 10），水溶性 68.84 ± 6.68%（n = 3）。这里采用克重更容易复核的 1.00% w/w 预配醋酸液，属于项目转换配方，不能声称与原文完全相同。")
    para(d, "Zullo 等纯壳聚糖膜论文采用无甘油配方，偏硬，并没有提供袋体接缝或湿态承重数据。本流程的胶接、提手、目标尺寸和荷载梯度是新的工程试验设定，失败也是有效结果。")

    heading(d, "二 单袋与三袋投料清单")
    table(d, ["材料", "每张 15 × 15 cm 膜", "一袋两张膜及胶液", "三袋独立重复"], [
        ("干壳聚糖粉", "1.00 g", "袋身 2.00 g + 胶液 0.10 g = 2.10 g", "6.30 g"),
        ("甘油", "0.30 g", "0.60 g", "1.80 g"),
        ("预配 1.00% w/w 醋酸水溶液", "100.00 g", "袋身 200.00 g + 胶液 10.00 g = 210.00 g", "630.00 g"),
        ("其中纯醋酸 / 水", "1.00 g / 99.00 g", "2.10 g / 207.90 g", "6.30 g / 623.70 g"),
    ], [1.7, 1.35, 2.6, 1.35])
    para(d, "称量口径：各张膜的成膜混合物投入质量为 101.30 g；两张共 202.60 g。胶液由 0.10 g 壳聚糖 + 10.00 g 预配酸液组成，总量 10.10 g，单袋最多计划使用 5.00 g，其余称量回收或按实验室要求处理。称量总投入不等于干膜实测质量；醋酸、水和脱模损失均需单独记录。")
    para(d, "采购要求：优先选有产品合格证、脱乙酰度和粘度/分子量批次信息、SDS 的商品壳聚糖；同一组三袋使用同一批号。壳聚糖先按产品允许的条件低温干燥至恒重，再称取表中的干基质量；若只能按到货质量称量，须另测含水率并换算。1.00% w/w 醋酸水溶液应购买预配液或由有资质人员按质量配制和标定，不让未经培训者接触冰乙酸。材料全部按非食品接触样品管理。")

    heading(d, "三 设备与环境")
    para(d, "0.001 g 分析天平、两个内部尺寸 15 × 15 cm 的平底模具、水平仪、搅拌器、65°C 可控干燥箱、温湿度计、厚度计、量尺、裁刀、刷或移液器、砝码/称重袋、护目镜和适用手套。避免在厨房使用食用器皿；壳聚糖称量时降低扬尘，甲壳类过敏者不要参与。")

    heading(d, "四 从粉末到两张膜")
    table(d, ["序号", "操作", "必须记录"], [
        ("1", "每个模具分别称干壳聚糖 1.00 g、甘油 0.30 g、预配 1.00% w/w 醋酸水溶液 100.00 g。先使壳聚糖在酸液中充分分散溶解，再加入甘油；室温持续搅拌约 24 h，静置脱泡约 2 h。两张膜分开配制。", "原料批号、实际克重、温度、是否完全溶解、有无颗粒和气泡。"),
        ("2", "模具水平放置，将一整批 101.30 g 成膜液倒入一个 15 × 15 cm 模具。两模相同。65°C 初干 5 h；冷却后检查，仍黏或质量不稳定则按 1 h 增量延长并记录。", "实际入模质量、模具面积、干燥时长、脱模前后质量。"),
        ("3", "完整脱模。调湿条件设为 25 ± 2°C、相对湿度 50 ± 5% 下 48 h；若没有恒湿箱，记录真实温湿度，结果仅作内部比较。每张膜测中央及四角共五点厚度。", "两张膜各自的五点厚度、均值、外观、折叠或破裂位置。"),
    ], [0.5, 4.25, 2.25])
    para(d, "理论干聚合物加甘油面密度为 1.30 g ÷ 225 cm² = 5.78 mg/cm²；这只是投料几何值。厚度必须实测，不能从投料保证会得到论文的 49 μm。")

    heading(d, "五 裁切 封边和提手")
    para(d, "两张膜各裁 12.0 × 14.0 cm 作前后袋身。另从两张模具边料裁出四条 1.5 × 10.0 cm 提手条，以及四块 2 × 2 cm 补强片。两条提手条层叠成一根提手；共做两根。边料面积够用，但若膜边缘不完整，另制第三张膜并按同一配方记录投料。")
    para(d, "胶液：称壳聚糖 0.10 g 与预配酸液 10.00 g，搅拌到无可见颗粒。先用边料做 5 mm 搭接试条；若干燥后轻拉即剥离，停止制作承重袋并记录失败。试条通过后，两袋身的左右边及底边各做 5 mm 搭接，三边约 40 cm；目标胶液 4.00 g。提手两端分别置于袋口内侧，并覆 2 × 2 cm 补强片，四个连接点合计目标胶液 1.00 g。实际使用量通过胶液容器前后称重计算，不能只按刷涂次数估计。")
    para(d, "接缝与提手在平面夹持、室温通风条件下至少干燥 24 h，继续至不黏手、接缝质量稳定；再按同一调湿环境放置 48 h。裁 12 × 14 cm、三边各 5 mm 搭接后，名义袋内宽约 11.0 cm、内高约 13.5 cm；实际内尺寸以成品量尺为准。热封参数尚无文献证据，本流程不使用热封。")

    heading(d, "六 性能检验和判定")
    table(d, ["项目", "操作及数据", "结果写法"], [
        ("膜", "每片测五点厚度；折叠 10 次；剪试条做干态拉伸，记录最大载荷和断裂伸长。", "三袋共六片；报告单片与批间差异。"),
        ("接缝", "每袋先拉独立胶接试条；成袋后检查三边是否连续、有无剥离或漏口。", "注明破坏发生在膜、胶层还是补强片。"),
        ("干态提举", "空袋合格后逐级装 50、100、200、500 g 非尖锐物；每级静置 60 s，通过后提起放下 10 次。", "只报告实际通过的最高重量与失败点；此梯度不是任何袋类标准。"),
        ("潮湿复测", "另取样品在约 75% RH 放置 24 h；若不能完整提起即记失败。", "记录温湿度、增重、变形、接缝开裂及最高实际荷载。"),
    ], [1.2, 3.65, 2.15])
    para(d, "作为本轮研发目标，至少三只独立样品在干态 200 g、60 s 静载和随后 10 次提举中无破裂，且潮湿样品能维持结构完整；这是项目筛选目标，不是产品标准。若未达到，保留失败样品并改变单一变量重新制膜。不得用‘降解’‘食品级’或‘环保认证’描述未经对应检测的成品。")

    heading(d, "七 安全与异常处理")
    bullet(d, "1% 醋酸水溶液仍会刺激眼睛；干燥箱与热模具可能烫伤。操作前阅读商品壳聚糖、醋酸液和甘油的 SDS；溶液入眼立即用流动清水冲洗并按当地急救流程处理。")
    bullet(d, "若称量时出现甲壳类相关过敏症状、呼吸困难，立即离开暴露环境并求医。不要使用来源不明或未标明浓度的酸液。")
    bullet(d, "所有样品贴批次、配方和日期标签；不得装食物、饮料、药物，也不得让儿童接触化学实验材料。")

    heading(d, "八 操作记录模板")
    table(d, ["批次", "膜片实投 / g", "干膜质量 / g", "五点厚度 / μm", "接缝胶液 / g", "干态最高通过载荷 / g", "失效位置"], [
        ("A", "", "", "", "", "", ""),
        ("B", "", "", "", "", "", ""),
        ("C", "", "", "", "", "", ""),
    ], [0.45, 1.05, 0.9, 1.1, 0.95, 1.45, 1.1])
    heading(d, "主要依据")
    small(d, "Jiménez-Regalado 等. Preparation and Physicochemical Properties of Modified Corn Starch-Chitosan Biodegradable Films. Polymers, 2021. DOI: 10.3390/polym13244431，方法 2.2，表 1。纯壳聚糖组含商品壳聚糖和甘油；本文按质量法重新定义稀醋酸用量。")
    small(d, "Zullo R 等. New Active Biopolymers and Chitosan-Based Films from Non-Native Crayfish Shell. Molecules, 2026. DOI: 10.3390/molecules31162819，方法 3.4。用于 24 h 溶解、2 h 脱泡及独立批次参考；其原配方无甘油。")
    d.save(ROOT / "商品壳聚糖_小袋样品_操作流程.docx")


if __name__ == "__main__":
    lab_proposal()
    bought_chitosan_protocol()
